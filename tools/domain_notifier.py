#!/usr/bin/env python3
"""
Ежедневный контроль сроков продления доменов с уведомлениями в Telegram.

Поддерживаемые поля Excel:
- Домен
- Дата окончания
- Регистратор
- Статус / Решение / Брать / Хороший
- Комментарий / Примечание / Заметка

Значения статуса распознаются автоматически:
- «брать», «да», «хороший», «продлеваем» -> продлеваем;
- «нет», «не брать», «не продлеваем» -> не продлеваем;
- «50 на 50», «50/50», «под вопросом» -> нужно решить.

Также поддерживаются обычные Excel-примечания к ячейкам: их текст будет
добавлен к комментарию соответствующего домена.

Запуск:
    python domain_notifier.py

Проверка без отправки в Telegram:
    python domain_notifier.py --dry-run

Запуск для произвольной даты:
    python domain_notifier.py --dry-run --date 2026-08-03

Переменные окружения:
    TELEGRAM_BOT_TOKEN       токен Telegram-бота (обязательно)
    TELEGRAM_CHANNEL         канал/чат, например @InfoDomVP
    SEND_STARTUP_MESSAGE     1 — отправлять сообщение о запуске, 0 — не отправлять
    SEND_OK_MESSAGE          1 — отправлять «всё в порядке», 0 — молчать
    EXCLUDE_DO_NOT_RENEW     1 — не уведомлять о доменах «не продлеваем»
    ALERT_BEFORE_DAYS        за сколько дней начинать уведомления (по умолчанию 60)

Cron, каждый день в 09:00:
    0 9 * * * /usr/bin/python3 /path/to/domain_notifier.py
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
import time
from collections import Counter
from datetime import date, datetime
from html import escape
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import requests
from openpyxl.utils.datetime import from_excel


# ─────────────────────────────────────────────
# НАСТРОЙКИ
# ─────────────────────────────────────────────
BASE_DIR = Path(__file__).resolve().parent
XLSX_FILE = BASE_DIR / "domains_by_month.xlsx"
SHEET_NAME = "Все домены (по дате)"
LOG_FILE = BASE_DIR / "domain_notifier.log"

BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
CHANNEL = os.environ.get("TELEGRAM_CHANNEL", "").strip()

ALERT_BEFORE_DAYS = int(os.environ.get("ALERT_BEFORE_DAYS", "60"))
SEND_STARTUP_MESSAGE = os.environ.get("SEND_STARTUP_MESSAGE", "0") == "1"
SEND_OK_MESSAGE = os.environ.get("SEND_OK_MESSAGE", "1") == "1"
EXCLUDE_DO_NOT_RENEW = os.environ.get("EXCLUDE_DO_NOT_RENEW", "0") == "1"

TELEGRAM_MAX_LENGTH = 4096
# Оставляем запас на служебные строки и номер части.
TELEGRAM_SAFE_LENGTH = 3700
MAX_COMMENT_LENGTH = 900
MAX_RETRIES = 3
RETRY_DELAY_SECONDS = 5

REGISTRAR_LINKS = {
    "reg.ru": "https://www.reg.ru/",
    "backorder": "https://backorder.ru/",
    "webnames": "https://www.webnames.ru/",
    "active.domains": "https://my.active.domains/",
    "i7": "https://my.i7.ru/billmgr",
    "hoster.by": "https://cp.hoster.by/personal/profile",
}

# Возможные названия столбцов. Регистр и лишние знаки не важны.
COLUMN_ALIASES = {
    "domain": {
        "домен",
        "domain",
        "доменное имя",
        "имя домена",
    },
    "expiry": {
        "дата окончания",
        "окончание",
        "истекает",
        "дата истечения",
        "expiry",
        "expiry date",
    },
    "registrar": {
        "регистратор",
        "registrar",
    },
    "status": {
        "статус",
        "решение",
        "брать",
        "продлевать",
        "продлеваем",
        "оставляем",
        "хороший",
        "оценка",
        "продление",
        "решение по домену",
        "хороший не продлеваем",
        "брать нет",
        "брать нет 50 на 50",
    },
    "comment": {
        "комментарий",
        "комментарии",
        "коммент",
        "примечание",
        "примечания",
        "заметка",
        "заметки",
        "описание",
        "comment",
        "comments",
        "note",
        "notes",
        "комментарий к домену",
        "комментарии к домену",
    },
}

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)s  %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ─────────────────────────────────────────────
def normalize_header(value: Any) -> str:
    """Нормализует название столбца для сопоставления с алиасами."""
    if value is None:
        return ""
    text = str(value).strip().lower().replace("ё", "е")
    text = re.sub(r"[^a-zа-я0-9]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def normalize_text(value: Any) -> str:
    """Преобразует значение ячейки в аккуратную строку."""
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def parse_expiry(value: Any, workbook_epoch: datetime) -> date:
    """Преобразует Excel-значение или строку в дату."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value, epoch=workbook_epoch)
        return converted.date() if isinstance(converted, datetime) else converted

    text = normalize_text(value)
    if not text:
        raise ValueError("пустая дата")

    formats = (
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    # Иногда Excel/LibreOffice сохраняет ISO-дату с часовым поясом или долями секунд.
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError as exc:
        raise ValueError(f"неподдерживаемый формат даты: {text}") from exc


def find_header(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[int, dict[str, int]]:
    """Находит строку заголовков и номера нужных столбцов."""
    normalized_aliases = {
        field: {normalize_header(alias) for alias in aliases}
        for field, aliases in COLUMN_ALIASES.items()
    }

    max_scan_rows = min(ws.max_row, 25)
    for row_number in range(1, max_scan_rows + 1):
        mapping: dict[str, int] = {}
        for column_number in range(1, ws.max_column + 1):
            header = normalize_header(ws.cell(row=row_number, column=column_number).value)
            if not header:
                continue
            for field, aliases in normalized_aliases.items():
                if header in aliases and field not in mapping:
                    mapping[field] = column_number

        if "domain" in mapping and "expiry" in mapping:
            return row_number, mapping

    raise ValueError(
        "Не найдена строка заголовков. Обязательны столбцы "
        "«Домен» и «Дата окончания»."
    )


def classify_status(value: Any) -> tuple[str, str]:
    """Возвращает внутренний код и отображаемый текст статуса."""
    if isinstance(value, bool):
        return ("keep", "✅ Брать / продлеваем") if value else ("drop", "❌ Не продлеваем")

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == 1:
            return "keep", "✅ Брать / продлеваем"
        if value == 0:
            return "drop", "❌ Не продлеваем"
        if value == 0.5 or value == 50:
            return "maybe", "🤔 50 на 50"

    raw = normalize_text(value)
    text = normalize_header(raw)
    compact = text.replace(" ", "")

    if not text:
        return "unknown", "⚪ Не указано"

    keep_values = {
        "брать",
        "да",
        "хороший",
        "хорошая",
        "хорошее",
        "продлеваем",
        "продлевать",
        "оставляем",
        "оставить",
        "нужен",
        "нужна",
        "keep",
        "good",
        "+",
    }
    drop_values = {
        "нет",
        "не брать",
        "не берем",
        "не берём",
        "не продлеваем",
        "не продлевать",
        "не оставляем",
        "не оставить",
        "плохой",
        "плохая",
        "drop",
        "delete",
        "no",
        "-",
    }
    maybe_values = {
        "50 на 50",
        "50 50",
        "50/50",
        "под вопросом",
        "сомнительно",
        "сомнение",
        "решить",
        "думаем",
        "maybe",
    }

    if text in keep_values:
        return "keep", "✅ Брать / продлеваем"
    if text in drop_values or text.startswith("не продлев") or text.startswith("не брать"):
        return "drop", "❌ Не продлеваем"
    if text in maybe_values or compact in {"50/50", "5050", "50на50"}:
        return "maybe", "🤔 50 на 50"

    return "other", f"ℹ️ {raw}"


def collect_excel_notes(cells: Iterable[openpyxl.cell.cell.Cell]) -> list[str]:
    """Собирает обычные Excel-примечания из ячеек строки."""
    notes: list[str] = []
    for cell in cells:
        if cell.comment and cell.comment.text:
            note = normalize_text(cell.comment.text)
            if note and note not in notes:
                notes.append(note)
    return notes


def merge_comments(column_comment: Any, excel_notes: Iterable[str]) -> str:
    """Объединяет комментарий из столбца и Excel-примечания."""
    parts: list[str] = []
    direct_comment = normalize_text(column_comment)
    if direct_comment:
        parts.append(direct_comment)

    for note in excel_notes:
        if note and note not in parts:
            parts.append(note)

    result = " / ".join(parts)
    if len(result) > MAX_COMMENT_LENGTH:
        result = result[: MAX_COMMENT_LENGTH - 1].rstrip() + "…"
    return result


# ─────────────────────────────────────────────
# ЗАГРУЗКА EXCEL
# ─────────────────────────────────────────────
def load_domains(xlsx_path: Path, sheet_name: str = SHEET_NAME) -> list[dict[str, Any]]:
    """Загружает домены, статусы и комментарии из Excel."""
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Файл не найден: {xlsx_path}")

    try:
        # read_only=False нужен для чтения Excel-примечаний (cell.comment).
        wb = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=True)
    except Exception as exc:
        raise RuntimeError(f"Ошибка открытия файла {xlsx_path}: {exc}") from exc

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Лист «{sheet_name}» не найден. Доступные листы: {', '.join(wb.sheetnames)}"
            )

        ws = wb[sheet_name]
        header_row, columns = find_header(ws)

        log.info(
            "Найдены столбцы: %s",
            ", ".join(f"{name}={number}" for name, number in columns.items()),
        )

        if "status" not in columns:
            log.warning(
                "Столбец статуса не найден. Будет показан статус «Не указано». "
                "Поддерживаемые заголовки: Статус, Решение, Брать, Хороший."
            )
        if "comment" not in columns:
            log.warning(
                "Столбец комментария не найден. Будут использованы Excel-примечания, если они есть."
            )

        domains: list[dict[str, Any]] = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            row_cells = [ws.cell(row=row_number, column=col) for col in range(1, ws.max_column + 1)]

            domain_value = ws.cell(row=row_number, column=columns["domain"]).value
            expiry_value = ws.cell(row=row_number, column=columns["expiry"]).value
            domain = normalize_text(domain_value)

            if not domain and expiry_value in (None, ""):
                continue
            if not domain:
                log.warning("Строка %s: пропущена — не указан домен", row_number)
                continue
            if expiry_value in (None, ""):
                log.warning("Строка %s: пропущена — не указана дата для %s", row_number, domain)
                continue

            try:
                expiry = parse_expiry(expiry_value, wb.epoch)
            except (ValueError, TypeError) as exc:
                log.warning(
                    "Строка %s: не удалось распознать дату %r для %s: %s",
                    row_number,
                    expiry_value,
                    domain,
                    exc,
                )
                continue

            registrar = "—"
            if "registrar" in columns:
                registrar = normalize_text(
                    ws.cell(row=row_number, column=columns["registrar"]).value
                ) or "—"

            excel_notes = collect_excel_notes(row_cells)
            status_value = (
                ws.cell(row=row_number, column=columns["status"]).value
                if "status" in columns
                else None
            )

            # Если отдельного статуса нет, но одно из Excel-примечаний само является
            # статусом («брать», «нет», «50/50»), используем его как статус.
            if status_value in (None, ""):
                for note in list(excel_notes):
                    status_code, _ = classify_status(note)
                    if status_code in {"keep", "drop", "maybe"}:
                        status_value = note
                        excel_notes.remove(note)
                        break

            status_code, status_label = classify_status(status_value)
            column_comment = (
                ws.cell(row=row_number, column=columns["comment"]).value
                if "comment" in columns
                else None
            )
            comment = merge_comments(column_comment, excel_notes)

            domains.append(
                {
                    "domain": domain,
                    "expiry": expiry,
                    "registrar": registrar,
                    "status_code": status_code,
                    "status": status_label,
                    "status_raw": normalize_text(status_value),
                    "comment": comment,
                    "row": row_number,
                }
            )

        log.info("Загружено доменов: %s", len(domains))
        status_counts = Counter(item["status_code"] for item in domains)
        log.info("Статусы: %s", dict(status_counts))
        return domains
    finally:
        wb.close()


# ─────────────────────────────────────────────
# TELEGRAM И ФОРМАТИРОВАНИЕ
# ─────────────────────────────────────────────
def days_label(number: int) -> str:
    """Склоняет слово «день»."""
    number = abs(number)
    if number % 10 == 1 and number % 100 != 11:
        return f"{number} день"
    if 2 <= number % 10 <= 4 and not (12 <= number % 100 <= 14):
        return f"{number} дня"
    return f"{number} дней"


def urgency_emoji(days_left: int) -> str:
    if days_left <= 0:
        return "🔴"
    if days_left <= 3:
        return "🔴"
    if days_left <= 7:
        return "🟠"
    if days_left <= 14:
        return "🟡"
    if days_left <= 30:
        return "🔵"
    return "⚪"


def expiry_heading(days_left: int) -> str:
    emoji = urgency_emoji(days_left)
    if days_left == 0:
        return f"{emoji} <b>Истекает сегодня:</b>"
    if days_left > 0:
        return f"{emoji} <b> {days_label(days_left)}:</b>"
    return f"{emoji} <b>Истёк {days_label(days_left)} назад:</b>"


def registrar_html(registrar: str) -> str:
    """Возвращает безопасный HTML регистратора; известные регистраторы остаются кликабельными."""
    if not registrar or registrar == "—":
        return "—"

    safe_registrar = escape(registrar)
    normalized = registrar.strip().lower()
    for key, url in REGISTRAR_LINKS.items():
        if key in normalized:
            # Ссылка сохраняется, но лишний значок 🔗 не показываем.
            return f'<a href="{escape(url, quote=True)}">{safe_registrar}</a>'
    return safe_registrar


def status_emoji(status_code: str) -> str:
    """Короткий итог по домену для компактного Telegram-уведомления."""
    return {
        "keep": "✅",
        "drop": "❌",
        "maybe": "🤔",
        "unknown": "⚪",
    }.get(status_code, "ℹ️")


def format_domain_line(alert: dict[str, Any]) -> str:
    """Форматирует один домен строго в одну строку."""
    safe_domain = escape(alert["domain"])
    registrar = registrar_html(alert["registrar"])
    decision = status_emoji(alert["status_code"])

    return (
        f"{expiry_heading(alert['days_left'])} "
        f"• <code>{safe_domain}</code> | "
        f"📅 {alert['expiry'].strftime('%d.%m.%Y')} | "
        f"🏢 {registrar} {decision}"
    )


def build_alert_messages(alerts: list[dict[str, Any]], check_date: date) -> list[str]:
    """Формирует компактные сообщения: один домен = одна строка."""
    if not alerts:
        return []

    sorted_alerts = sorted(alerts, key=lambda item: (item["days_left"], item["domain"].lower()))
    messages: list[str] = []
    current_lines: list[str] = []

    for alert in sorted_alerts:
        line = format_domain_line(alert)
        candidate = "\n".join([*current_lines, line])

        if current_lines and len(candidate) > TELEGRAM_SAFE_LENGTH:
            messages.append("\n".join(current_lines))
            current_lines = [line]
        else:
            current_lines.append(line)

    if current_lines:
        messages.append("\n".join(current_lines))

    for index, message in enumerate(messages, start=1):
        if len(message) > TELEGRAM_MAX_LENGTH:
            raise ValueError(
                f"Сообщение {index} имеет длину {len(message)} символов и превышает лимит Telegram"
            )

    return messages


def build_ok_message(check_date: date, domains: list[dict[str, Any]]) -> str:
    status_counts = Counter(item["status_code"] for item in domains)
    return (
        "✅ <b>Мониторинг доменов</b>\n"
        f"🗓 {check_date.strftime('%d.%m.%Y')}\n\n"
        f"Доменов для уведомления в ближайшие {ALERT_BEFORE_DAYS} дней нет.\n"
        f"📊 Всего в базе: {len(domains)}\n"
        f"✅ Продлеваем: {status_counts.get('keep', 0)}\n"
        f"❌ Не продлеваем: {status_counts.get('drop', 0)}\n"
        f"🤔 50 на 50: {status_counts.get('maybe', 0)}\n"
        f"⚪ Без решения: {status_counts.get('unknown', 0)}"
    )


def build_startup_message(check_date: date, domains: list[dict[str, Any]]) -> str:
    return (
        "🚀 <b>Мониторинг доменов запущен</b>\n"
        f"🗓 {check_date.strftime('%d.%m.%Y')}\n"
        f"📊 Всего доменов в базе: {len(domains)}\n"
        f"📋 Уведомления: от {ALERT_BEFORE_DAYS} дней до дня окончания включительно"
    )


def send_telegram(text: str) -> bool:
    """Отправляет одно HTML-сообщение в Telegram."""
    if not BOT_TOKEN:
        log.error("Не задана переменная окружения TELEGRAM_BOT_TOKEN")
        return False
    if not CHANNEL:
        log.error("Не задана переменная окружения TELEGRAM_CHANNEL")
        return False
    if len(text) > TELEGRAM_MAX_LENGTH:
        log.error("Сообщение слишком длинное: %s символов", len(text))
        return False

    url = f"https://api.telegram.org/bot{BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id": CHANNEL,
        "text": text,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }

    try:
        response = requests.post(url, json=payload, timeout=30)
        response.raise_for_status()
        try:
            result = response.json()
        except ValueError:
            log.error("Telegram вернул не-JSON ответ: %s", response.text[:500])
            return False

        if not result.get("ok"):
            log.error("Telegram вернул ошибку: %s", result)
            return False
        return True
    except requests.exceptions.Timeout:
        log.error("Таймаут при отправке в Telegram")
    except requests.exceptions.RequestException as exc:
        log.error("Ошибка отправки в Telegram: %s", exc)
    return False


def send_with_retries(text: str, label: str) -> bool:
    for attempt in range(1, MAX_RETRIES + 1):
        log.info("Отправка %s, попытка %s/%s", label, attempt, MAX_RETRIES)
        if send_telegram(text):
            log.info("%s успешно отправлено", label)
            return True
        if attempt < MAX_RETRIES:
            log.warning("Ошибка отправки; повтор через %s секунд", RETRY_DELAY_SECONDS)
            time.sleep(RETRY_DELAY_SECONDS)
    log.error("Не удалось отправить %s после %s попыток", label, MAX_RETRIES)
    return False


# ─────────────────────────────────────────────
# ОСНОВНАЯ ЛОГИКА
# ─────────────────────────────────────────────
def parse_check_date(value: str | None) -> date:
    if not value:
        return date.today()
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Дата должна быть в формате YYYY-MM-DD") from exc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Мониторинг сроков доменов")
    parser.add_argument(
        "--file",
        type=Path,
        default=XLSX_FILE,
        help=f"Путь к Excel-файлу (по умолчанию: {XLSX_FILE})",
    )
    parser.add_argument(
        "--sheet",
        default=SHEET_NAME,
        help=f"Имя листа (по умолчанию: {SHEET_NAME})",
    )
    parser.add_argument(
        "--date",
        dest="check_date",
        help="Дата проверки YYYY-MM-DD; полезно для тестирования",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Ничего не отправлять, а вывести сообщения в консоль",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    today = parse_check_date(args.check_date)

    log.info("=" * 60)
    log.info("Запуск проверки. Дата: %s", today)
    log.info("Уведомлять от %s дней до дня окончания", ALERT_BEFORE_DAYS)
    log.info("Файл: %s", args.file)
    log.info("=" * 60)

    try:
        domains = load_domains(args.file, args.sheet)
    except (FileNotFoundError, RuntimeError, ValueError) as exc:
        log.error("%s", exc)
        return 1

    if not domains:
        log.warning("Нет доменов для проверки")
        return 0

    alerts: list[dict[str, Any]] = []
    for domain_data in domains:
        days_left = (domain_data["expiry"] - today).days
        if not 0 <= days_left <= ALERT_BEFORE_DAYS:
            continue
        if EXCLUDE_DO_NOT_RENEW and domain_data["status_code"] == "drop":
            log.info(
                "Пропущен домен «не продлеваем»: %s (%s)",
                domain_data["domain"],
                domain_data["expiry"],
            )
            continue

        alert = {**domain_data, "days_left": days_left}
        alerts.append(alert)
        log.info(
            "⚠️ %s — %s дн., статус: %s, комментарий: %s",
            domain_data["domain"],
            days_left,
            domain_data["status"],
            domain_data["comment"] or "—",
        )

    outgoing: list[tuple[str, str]] = []
    if SEND_STARTUP_MESSAGE:
        outgoing.append(("сообщение о запуске", build_startup_message(today, domains)))

    if alerts:
        messages = build_alert_messages(alerts, today)
        for index, message in enumerate(messages, start=1):
            outgoing.append((f"уведомление {index}/{len(messages)}", message))
        log.info("Найдено доменов для уведомления: %s", len(alerts))
    else:
        log.info("Доменов для уведомления сегодня нет")
        if SEND_OK_MESSAGE:
            outgoing.append(("сообщение «всё в порядке»", build_ok_message(today, domains)))

    if args.dry_run:
        print("\n" + "=" * 80)
        print("DRY RUN: сообщения не отправляются")
        print("=" * 80)
        for label, message in outgoing:
            print(f"\n--- {label} ({len(message)} символов) ---\n")
            print(message)
        if not outgoing:
            print("Нет сообщений для вывода.")
        return 0

    if not outgoing:
        log.info("Сообщений для отправки нет")
        return 0

    if not BOT_TOKEN:
        log.error(
            "TELEGRAM_BOT_TOKEN не задан. Установите переменную окружения "
            "или используйте --dry-run для проверки."
        )
        return 1

    all_sent = True
    for label, message in outgoing:
        if not send_with_retries(message, label):
            all_sent = False

    log.info("Проверка завершена")
    return 0 if all_sent else 2


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        log.info("Скрипт остановлен пользователем")
        raise SystemExit(130)
    except Exception as exc:
        log.error("Критическая ошибка: %s", exc, exc_info=True)
        raise SystemExit(1)
